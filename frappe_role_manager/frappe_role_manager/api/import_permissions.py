"""IMPORT API - Import Permissions from JSON and Excel.

Allows importing permissions from:
- JSON format
- Excel format

Features:
- Preview changes before applying
- Validation of data
- Conflict detection
- Backup before import
"""

import json

import frappe
from frappe import _
from frappe.utils import cint

from frappe_role_manager.frappe_role_manager.doctype.bulk_operation_log.bulk_operation_log import (
	create_operation_log,
)


@frappe.whitelist()
def preview_import(file_url=None, json_data=None):
	"""Preview import changes without applying them.

	Args:
	    file_url: URL of the uploaded Excel/JSON file
	    json_data: Direct JSON data string

	Returns:
	    dict: Preview of changes with validation results
	"""
	# Parse input data
	if json_data:
		if isinstance(json_data, str):
			import_data = json.loads(json_data)
		else:
			import_data = json_data
	elif file_url:
		import_data = _parse_file(file_url)
	else:
		frappe.throw(_("Either file_url or json_data is required"))

	# Validate and preview
	preview = {
		"role_permissions": {"add": [], "update": [], "errors": []},
		"user_permissions": {"add": [], "update": [], "remove": [], "errors": []},
		"role_assignments": {"add": [], "remove": [], "errors": []},
		"validation": {"errors": [], "warnings": []},
	}

	# Validate role permissions
	for perm in import_data.get("role_permissions", []):
		result = _validate_role_permission(perm)
		if result["valid"]:
			if result["exists"]:
				preview["role_permissions"]["update"].append(perm)
			else:
				preview["role_permissions"]["add"].append(perm)
		else:
			preview["role_permissions"]["errors"].append({"data": perm, "error": result["error"]})
			preview["validation"]["errors"].append(result["error"])

	# Validate user permissions
	for perm in import_data.get("user_permissions", []):
		result = _validate_user_permission(perm)
		if result["valid"]:
			action = perm.get("action", "add")
			if action == "remove":
				if result["exists"]:
					preview["user_permissions"]["remove"].append(perm)
				else:
					preview["validation"]["warnings"].append(f"User permission for {perm.get('user')} on {perm.get('allow')}={perm.get('for_value')} doesn't exist, skipping removal")
			elif result["exists"]:
				preview["user_permissions"]["update"].append(perm)
			else:
				preview["user_permissions"]["add"].append(perm)
		else:
			preview["user_permissions"]["errors"].append({"data": perm, "error": result["error"]})
			preview["validation"]["errors"].append(result["error"])

	# Validate role assignments
	for assignment in import_data.get("role_assignments", []):
		user = assignment.get("user")
		roles = assignment.get("roles", [])

		# Handle flat format (user, role, action)
		if "role" in assignment:
			roles = [assignment.get("role")]
			action = assignment.get("action", "add")

			result = _validate_role_assignment(user, assignment.get("role"))
			if result["valid"]:
				if action == "remove":
					if result["exists"]:
						preview["role_assignments"]["remove"].append({"user": user, "role": assignment.get("role")})
					else:
						preview["validation"]["warnings"].append(f"User {user} doesn't have role {assignment.get('role')}, skipping removal")
				elif not result["exists"]:
					preview["role_assignments"]["add"].append({"user": user, "role": assignment.get("role")})
			else:
				preview["role_assignments"]["errors"].append({"data": assignment, "error": result["error"]})
				preview["validation"]["errors"].append(result["error"])
		else:
			# Handle grouped format (user, roles[])
			for role in roles:
				result = _validate_role_assignment(user, role)
				if result["valid"]:
					if not result["exists"]:
						preview["role_assignments"]["add"].append({"user": user, "role": role})
				else:
					preview["role_assignments"]["errors"].append({"data": {"user": user, "role": role}, "error": result["error"]})
					preview["validation"]["errors"].append(result["error"])

	# Summary
	preview["summary"] = {
		"role_permissions_add": len(preview["role_permissions"]["add"]),
		"role_permissions_update": len(preview["role_permissions"]["update"]),
		"user_permissions_add": len(preview["user_permissions"]["add"]),
		"user_permissions_update": len(preview["user_permissions"]["update"]),
		"user_permissions_remove": len(preview["user_permissions"]["remove"]),
		"role_assignments_add": len(preview["role_assignments"]["add"]),
		"role_assignments_remove": len(preview["role_assignments"]["remove"]),
		"total_errors": len(preview["validation"]["errors"]),
		"total_warnings": len(preview["validation"]["warnings"]),
	}

	preview["can_proceed"] = len(preview["validation"]["errors"]) == 0

	return preview


@frappe.whitelist()
def execute_import(file_url=None, json_data=None, create_backup=True, skip_errors=False):
	"""Execute the import operation.

	Args:
	    file_url: URL of the uploaded Excel/JSON file
	    json_data: Direct JSON data string
	    create_backup: Whether to create a backup before importing
	    skip_errors: Whether to continue on errors

	Returns:
	    dict: Results of the import operation
	"""
	# Parse input data
	if json_data:
		if isinstance(json_data, str):
			import_data = json.loads(json_data)
		else:
			import_data = json_data
	elif file_url:
		import_data = _parse_file(file_url)
	else:
		frappe.throw(_("Either file_url or json_data is required"))

	create_backup = _to_bool(create_backup)
	skip_errors = _to_bool(skip_errors)

	# Create operation log
	log = create_operation_log("Import Permissions", {"file_url": file_url, "skip_errors": skip_errors})
	log.mark_started()

	results = {
		"role_permissions": {"added": 0, "updated": 0, "errors": []},
		"user_permissions": {"added": 0, "updated": 0, "removed": 0, "errors": []},
		"role_assignments": {"added": 0, "removed": 0, "errors": []},
	}

	try:
		# Create backup if requested
		if create_backup:
			_create_full_backup()

		# Import role permissions
		for perm in import_data.get("role_permissions", []):
			try:
				result = _import_role_permission(perm)
				if result == "added":
					results["role_permissions"]["added"] += 1
				elif result == "updated":
					results["role_permissions"]["updated"] += 1
			except Exception as e:
				error_msg = f"Role permission {perm.get('role')}/{perm.get('doctype')}: {e!s}"
				results["role_permissions"]["errors"].append(error_msg)
				if not skip_errors:
					raise

		# Import user permissions
		for perm in import_data.get("user_permissions", []):
			try:
				action = perm.get("action", "add")
				if action == "remove":
					result = _remove_user_permission(perm)
					if result:
						results["user_permissions"]["removed"] += 1
				else:
					result = _import_user_permission(perm)
					if result == "added":
						results["user_permissions"]["added"] += 1
					elif result == "updated":
						results["user_permissions"]["updated"] += 1
			except Exception as e:
				error_msg = f"User permission {perm.get('user')}/{perm.get('allow')}: {e!s}"
				results["user_permissions"]["errors"].append(error_msg)
				if not skip_errors:
					raise

		# Import role assignments
		for assignment in import_data.get("role_assignments", []):
			user = assignment.get("user")

			# Handle flat format
			if "role" in assignment:
				action = assignment.get("action", "add")
				role = assignment.get("role")
				try:
					if action == "remove":
						result = _remove_role_assignment(user, role)
						if result:
							results["role_assignments"]["removed"] += 1
					else:
						result = _import_role_assignment(user, role)
						if result:
							results["role_assignments"]["added"] += 1
				except Exception as e:
					error_msg = f"Role assignment {user}/{role}: {e!s}"
					results["role_assignments"]["errors"].append(error_msg)
					if not skip_errors:
						raise
			else:
				# Handle grouped format
				for role in assignment.get("roles", []):
					try:
						result = _import_role_assignment(user, role)
						if result:
							results["role_assignments"]["added"] += 1
					except Exception as e:
						error_msg = f"Role assignment {user}/{role}: {e!s}"
						results["role_assignments"]["errors"].append(error_msg)
						if not skip_errors:
							raise

		frappe.db.commit()

		# Calculate totals
		total_success = (
			results["role_permissions"]["added"]
			+ results["role_permissions"]["updated"]
			+ results["user_permissions"]["added"]
			+ results["user_permissions"]["updated"]
			+ results["user_permissions"]["removed"]
			+ results["role_assignments"]["added"]
			+ results["role_assignments"]["removed"]
		)

		total_errors = (
			len(results["role_permissions"]["errors"])
			+ len(results["user_permissions"]["errors"])
			+ len(results["role_assignments"]["errors"])
		)

		log.mark_completed(success_count=total_success, failure_count=total_errors, details=results)

		return {"status": "success", "results": results, "log": log.name}

	except Exception as e:
		frappe.db.rollback()
		error_msg = str(e)
		log.mark_failed(error_msg)
		return {"status": "error", "message": error_msg, "results": results, "log": log.name}


def _parse_file(file_url):
	"""Parse uploaded file (Excel or JSON)."""
	file_doc = frappe.get_doc("File", {"file_url": file_url})
	file_path = file_doc.get_full_path()

	if file_path.endswith(".json"):
		with open(file_path) as f:
			return json.load(f)
	elif file_path.endswith((".xlsx", ".xls")):
		return _parse_excel(file_path)
	else:
		frappe.throw(_("Unsupported file format. Use .json, .xlsx, or .xls"))


def _parse_excel(file_path):
	"""Parse Excel file into import data format."""
	try:
		import openpyxl
	except ImportError:
		frappe.throw(_("openpyxl is required for Excel import. Install with: pip install openpyxl"))

	wb = openpyxl.load_workbook(file_path)
	import_data = {"role_permissions": [], "user_permissions": [], "role_assignments": []}

	# Parse Role Assignments sheet
	if "Role Assignments" in wb.sheetnames:
		ws = wb["Role Assignments"]
		headers = [cell.value for cell in ws[1]]
		for row in ws.iter_rows(min_row=2, values_only=True):
			if not row[0]:  # Skip empty rows
				continue
			row_data = dict(zip(headers, row))
			if row_data.get("User") and row_data.get("Role"):
				import_data["role_assignments"].append(
					{"user": row_data["User"], "role": row_data["Role"], "action": row_data.get("Action", "add")}
				)

	# Parse User Permissions sheet
	if "User Permissions" in wb.sheetnames:
		ws = wb["User Permissions"]
		headers = [cell.value for cell in ws[1]]
		for row in ws.iter_rows(min_row=2, values_only=True):
			if not row[0]:  # Skip empty rows
				continue
			row_data = dict(zip(headers, row))
			if row_data.get("User") and row_data.get("Allow DocType"):
				apply_to_all = row_data.get("Apply To All DocTypes", "Yes")
				if isinstance(apply_to_all, str):
					apply_to_all = apply_to_all.lower() in ("yes", "true", "1")

				import_data["user_permissions"].append(
					{
						"user": row_data["User"],
						"allow": row_data["Allow DocType"],
						"for_value": row_data.get("For Value"),
						"applicable_for": row_data.get("Applicable For"),
						"apply_to_all_doctypes": apply_to_all,
						"action": row_data.get("Action", "add"),
					}
				)

	# Parse Role Permissions sheet
	if "Role Permissions" in wb.sheetnames:
		ws = wb["Role Permissions"]
		headers = [cell.value for cell in ws[1]]
		for row in ws.iter_rows(min_row=2, values_only=True):
			if not row[0]:  # Skip empty rows
				continue
			row_data = dict(zip(headers, row))
			if row_data.get("Role") and row_data.get("DocType"):
				import_data["role_permissions"].append(
					{
						"role": row_data["Role"],
						"doctype": row_data["DocType"],
						"read": cint(row_data.get("Read", 0)),
						"write": cint(row_data.get("Write", 0)),
						"create": cint(row_data.get("Create", 0)),
						"delete": cint(row_data.get("Delete", 0)),
						"submit": cint(row_data.get("Submit", 0)),
						"cancel": cint(row_data.get("Cancel", 0)),
						"amend": cint(row_data.get("Amend", 0)),
						"report": cint(row_data.get("Report", 0)),
						"export": cint(row_data.get("Export", 0)),
						"import": cint(row_data.get("Import", 0)),
						"share": cint(row_data.get("Share", 0)),
						"print": cint(row_data.get("Print", 0)),
						"email": cint(row_data.get("Email", 0)),
					}
				)

	return import_data


def _validate_role_permission(perm):
	"""Validate a role permission entry."""
	role = perm.get("role")
	doctype = perm.get("doctype")

	if not role:
		return {"valid": False, "error": "Role is required"}
	if not doctype:
		return {"valid": False, "error": "DocType is required"}

	if not frappe.db.exists("Role", role):
		return {"valid": False, "error": f"Role '{role}' does not exist"}
	if not frappe.db.exists("DocType", doctype):
		return {"valid": False, "error": f"DocType '{doctype}' does not exist"}

	exists = frappe.db.exists("DocPerm", {"parent": doctype, "role": role, "permlevel": 0})

	return {"valid": True, "exists": bool(exists)}


def _validate_user_permission(perm):
	"""Validate a user permission entry."""
	user = perm.get("user")
	allow = perm.get("allow")
	for_value = perm.get("for_value")

	if not user:
		return {"valid": False, "error": "User is required"}
	if not allow:
		return {"valid": False, "error": "Allow DocType is required"}
	if not for_value:
		return {"valid": False, "error": "For Value is required"}

	if not frappe.db.exists("User", user):
		return {"valid": False, "error": f"User '{user}' does not exist"}
	if not frappe.db.exists("DocType", allow):
		return {"valid": False, "error": f"DocType '{allow}' does not exist"}

	# Check if the for_value exists
	if not frappe.db.exists(allow, for_value):
		return {"valid": False, "error": f"'{for_value}' does not exist in DocType '{allow}'"}

	exists = frappe.db.exists("User Permission", {"user": user, "allow": allow, "for_value": for_value})

	return {"valid": True, "exists": bool(exists)}


def _validate_role_assignment(user, role):
	"""Validate a role assignment entry."""
	if not user:
		return {"valid": False, "error": "User is required"}
	if not role:
		return {"valid": False, "error": "Role is required"}

	if not frappe.db.exists("User", user):
		return {"valid": False, "error": f"User '{user}' does not exist"}
	if not frappe.db.exists("Role", role):
		return {"valid": False, "error": f"Role '{role}' does not exist"}

	exists = frappe.db.exists("Has Role", {"parent": user, "parenttype": "User", "role": role})

	return {"valid": True, "exists": bool(exists)}


def _import_role_permission(perm):
	"""Import a single role permission."""
	doctype = perm.get("doctype")
	role = perm.get("role")

	existing = frappe.db.exists("DocPerm", {"parent": doctype, "role": role, "permlevel": 0})

	if existing:
		# Update existing
		doc = frappe.get_doc("DocPerm", existing)
		for field in ["read", "write", "create", "delete", "submit", "cancel", "amend", "report", "export", "import", "share", "print", "email"]:
			if field in perm:
				setattr(doc, field, perm[field])
		doc.save(ignore_permissions=True)
		return "updated"
	else:
		# Add new via DocType
		dt = frappe.get_doc("DocType", doctype)
		dt.append(
			"permissions",
			{
				"role": role,
				"permlevel": 0,
				"read": perm.get("read", 0),
				"write": perm.get("write", 0),
				"create": perm.get("create", 0),
				"delete": perm.get("delete", 0),
				"submit": perm.get("submit", 0),
				"cancel": perm.get("cancel", 0),
				"amend": perm.get("amend", 0),
				"report": perm.get("report", 0),
				"export": perm.get("export", 0),
				"import": perm.get("import", 0),
				"share": perm.get("share", 0),
				"print": perm.get("print", 0),
				"email": perm.get("email", 0),
			},
		)
		dt.save(ignore_permissions=True)
		return "added"


def _import_user_permission(perm):
	"""Import a single user permission."""
	user = perm.get("user")
	allow = perm.get("allow")
	for_value = perm.get("for_value")

	existing = frappe.db.exists("User Permission", {"user": user, "allow": allow, "for_value": for_value})

	if existing:
		# Update existing
		doc = frappe.get_doc("User Permission", existing)
		if "applicable_for" in perm:
			doc.applicable_for = perm["applicable_for"]
		if "apply_to_all_doctypes" in perm:
			doc.apply_to_all_doctypes = perm["apply_to_all_doctypes"]
		doc.save(ignore_permissions=True)
		return "updated"
	else:
		# Create new
		doc = frappe.new_doc("User Permission")
		doc.user = user
		doc.allow = allow
		doc.for_value = for_value
		doc.applicable_for = perm.get("applicable_for")
		doc.apply_to_all_doctypes = perm.get("apply_to_all_doctypes", 1)
		doc.insert(ignore_permissions=True)
		return "added"


def _remove_user_permission(perm):
	"""Remove a user permission."""
	user = perm.get("user")
	allow = perm.get("allow")
	for_value = perm.get("for_value")

	existing = frappe.db.exists("User Permission", {"user": user, "allow": allow, "for_value": for_value})

	if existing:
		frappe.delete_doc("User Permission", existing, ignore_permissions=True)
		return True
	return False


def _import_role_assignment(user, role):
	"""Import a single role assignment."""
	exists = frappe.db.exists("Has Role", {"parent": user, "parenttype": "User", "role": role})

	if not exists:
		user_doc = frappe.get_doc("User", user)
		user_doc.append("roles", {"role": role})
		user_doc.save(ignore_permissions=True)
		return True
	return False


def _remove_role_assignment(user, role):
	"""Remove a role assignment from a user."""
	user_doc = frappe.get_doc("User", user)
	roles_to_keep = [r for r in user_doc.roles if r.role != role]

	if len(roles_to_keep) < len(user_doc.roles):
		user_doc.roles = []
		for r in roles_to_keep:
			user_doc.append("roles", {"role": r.role})
		user_doc.save(ignore_permissions=True)
		return True
	return False


def _create_full_backup():
	"""Create a full backup of all permissions."""
	from frappe_role_manager.frappe_role_manager.api.export import export_permissions

	result = export_permissions(export_roles=True, export_user_permissions=True, export_role_assignments=True, format="json")

	# Update snapshot description
	if result.get("snapshot"):
		snapshot = frappe.get_doc("Permission Snapshot", result["snapshot"])
		snapshot.snapshot_type = "Full Backup"
		snapshot.description = "Automatic backup before import operation"
		snapshot.save(ignore_permissions=True)

	return result.get("snapshot")


def _to_bool(value):
	"""Convert various truthy values to boolean."""
	if isinstance(value, bool):
		return value
	if isinstance(value, str):
		return value.lower() in ("true", "1", "yes")
	return bool(value)
