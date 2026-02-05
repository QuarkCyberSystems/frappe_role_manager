"""EXPORT API - Export Permissions to JSON and Excel.

Allows exporting permissions to:
- JSON format (complete data, good for version control)
- Excel format (user-friendly, editable)
"""

import json

import frappe
from frappe import _
from frappe.utils import now_datetime


@frappe.whitelist()
def export_permissions(
	users=None,
	roles=None,
	doctypes=None,
	export_roles=True,
	export_user_permissions=True,
	export_role_assignments=True,
	export_role_profiles=False,
	format="json",
):
	"""Export permissions based on filters.

	Args:
	    users: List of users to export (optional, exports all if None)
	    roles: List of roles to export (optional, exports all if None)
	    doctypes: List of doctypes to export permissions for (optional)
	    export_roles: Whether to export role definitions/DocPerms
	    export_user_permissions: Whether to export user permission records
	    export_role_assignments: Whether to export user-role assignments
	    export_role_profiles: Whether to export role profiles
	    format: 'json' or 'excel'

	Returns:
	    dict: Export data or file URL
	"""
	# Parse JSON strings if needed
	if isinstance(users, str) and users:
		users = json.loads(users)
	if isinstance(roles, str) and roles:
		roles = json.loads(roles)
	if isinstance(doctypes, str) and doctypes:
		doctypes = json.loads(doctypes)

	# Convert string booleans
	export_roles = _to_bool(export_roles)
	export_user_permissions = _to_bool(export_user_permissions)
	export_role_assignments = _to_bool(export_role_assignments)
	export_role_profiles = _to_bool(export_role_profiles)

	export_data = {
		"meta": {
			"exported_at": str(now_datetime()),
			"exported_by": frappe.session.user,
			"version": "1.0",
			"filters": {
				"users": users,
				"roles": roles,
				"doctypes": doctypes,
			},
		},
		"role_permissions": [],
		"user_permissions": [],
		"role_assignments": [],
		"role_profiles": [],
	}

	# Export role permissions (DocPerm)
	if export_roles:
		export_data["role_permissions"] = _get_role_permissions(roles, doctypes)

	# Export user permissions
	if export_user_permissions:
		export_data["user_permissions"] = _get_user_permissions(users)

	# Export role assignments
	if export_role_assignments:
		export_data["role_assignments"] = _get_role_assignments(users, roles)

	# Export role profiles
	if export_role_profiles:
		export_data["role_profiles"] = _get_role_profiles()

	# Create snapshot record
	snapshot = _create_export_snapshot(export_data)

	if format == "excel":
		file_url = _generate_excel(export_data)
		return {"status": "success", "format": "excel", "file_url": file_url, "snapshot": snapshot}
	else:
		return {"status": "success", "format": "json", "data": export_data, "snapshot": snapshot}


@frappe.whitelist()
def download_template(template_type="full"):
	"""Download a blank Excel template for importing permissions.

	Args:
	    template_type: 'full', 'roles', 'user_permissions', 'role_assignments'

	Returns:
	    dict: File URL for the template
	"""
	try:
		import openpyxl
		from openpyxl.styles import Font, PatternFill
	except ImportError:
		frappe.throw(_("openpyxl is required for Excel export. Install with: pip install openpyxl"))

	wb = openpyxl.Workbook()

	# Header style
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

	if template_type in ("full", "role_assignments"):
		# Role Assignments sheet
		ws = wb.active
		ws.title = "Role Assignments"
		headers = ["User", "Role", "Action"]
		for col, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col, value=header)
			cell.font = header_font
			cell.fill = header_fill
		# Add sample row
		ws.cell(row=2, column=1, value="user@example.com")
		ws.cell(row=2, column=2, value="Sales User")
		ws.cell(row=2, column=3, value="add")
		# Add instructions
		ws.cell(row=4, column=1, value="Instructions:")
		ws.cell(row=5, column=1, value="- User: Enter the user email")
		ws.cell(row=6, column=1, value="- Role: Enter the exact role name")
		ws.cell(row=7, column=1, value="- Action: 'add' to assign, 'remove' to unassign")

	if template_type in ("full", "user_permissions"):
		# User Permissions sheet
		if template_type == "full":
			ws = wb.create_sheet("User Permissions")
		else:
			ws = wb.active
			ws.title = "User Permissions"
		headers = ["User", "Allow DocType", "For Value", "Apply To All DocTypes", "Action"]
		for col, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col, value=header)
			cell.font = header_font
			cell.fill = header_fill
		# Add sample row
		ws.cell(row=2, column=1, value="user@example.com")
		ws.cell(row=2, column=2, value="Company")
		ws.cell(row=2, column=3, value="My Company")
		ws.cell(row=2, column=4, value="Yes")
		ws.cell(row=2, column=5, value="add")
		# Add instructions
		ws.cell(row=4, column=1, value="Instructions:")
		ws.cell(row=5, column=1, value="- User: Enter the user email")
		ws.cell(row=6, column=1, value="- Allow DocType: The DocType to restrict (e.g., Company, Territory)")
		ws.cell(row=7, column=1, value="- For Value: The specific value to allow")
		ws.cell(row=8, column=1, value="- Apply To All DocTypes: 'Yes' or 'No'")
		ws.cell(row=9, column=1, value="- Action: 'add' to create, 'remove' to delete")

	if template_type in ("full", "roles"):
		# Role Permissions sheet
		if template_type == "full":
			ws = wb.create_sheet("Role Permissions")
		else:
			ws = wb.active
			ws.title = "Role Permissions"
		headers = ["Role", "DocType", "Read", "Write", "Create", "Delete", "Submit", "Cancel", "Amend", "Report", "Export", "Import", "Share", "Print", "Email"]
		for col, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col, value=header)
			cell.font = header_font
			cell.fill = header_fill
		# Add sample row
		sample_data = ["Sales User", "Customer", 1, 1, 1, 0, 0, 0, 0, 1, 1, 0, 1, 1, 1]
		for col, value in enumerate(sample_data, 1):
			ws.cell(row=2, column=col, value=value)
		# Add instructions
		ws.cell(row=4, column=1, value="Instructions:")
		ws.cell(row=5, column=1, value="- Role: Enter the exact role name")
		ws.cell(row=6, column=1, value="- DocType: Enter the exact DocType name")
		ws.cell(row=7, column=1, value="- Permission columns: Use 1 for Yes, 0 for No")

	# Save to file
	from io import BytesIO

	file_buffer = BytesIO()
	wb.save(file_buffer)
	file_buffer.seek(0)

	# Create file in Frappe
	filename = f"permission_template_{template_type}.xlsx"
	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": filename,
			"content": file_buffer.getvalue(),
			"is_private": 0,
		}
	)
	file_doc.insert(ignore_permissions=True)

	return {"status": "success", "file_url": file_doc.file_url, "filename": filename}


def _get_role_permissions(roles=None, doctypes=None):
	"""Get DocPerm records based on filters."""
	filters = {"permlevel": 0}

	if roles:
		filters["role"] = ["in", roles]

	if doctypes:
		filters["parent"] = ["in", doctypes]

	docperms = frappe.get_all(
		"DocPerm",
		filters=filters,
		fields=[
			"parent as doctype",
			"role",
			"permlevel",
			"read",
			"write",
			"create",
			"delete",
			"submit",
			"cancel",
			"amend",
			"report",
			"export",
			"import",
			"share",
			"print",
			"email",
			"if_owner",
		],
		order_by="parent, role",
	)

	return docperms


def _get_user_permissions(users=None):
	"""Get User Permission records based on filters."""
	filters = {}

	if users:
		filters["user"] = ["in", users]

	user_perms = frappe.get_all(
		"User Permission",
		filters=filters,
		fields=["user", "allow", "for_value", "applicable_for", "apply_to_all_doctypes"],
		order_by="user, allow",
	)

	return user_perms


def _get_role_assignments(users=None, roles=None):
	"""Get user-role assignments based on filters."""
	filters = {"parenttype": "User"}

	if roles:
		filters["role"] = ["in", roles]

	if users:
		filters["parent"] = ["in", users]

	has_roles = frappe.get_all(
		"Has Role",
		filters=filters,
		fields=["parent as user", "role"],
		order_by="parent, role",
	)

	# Group by user
	grouped = {}
	for hr in has_roles:
		if hr.user not in grouped:
			grouped[hr.user] = {"user": hr.user, "roles": []}
		grouped[hr.user]["roles"].append(hr.role)

	return list(grouped.values())


def _get_role_profiles():
	"""Get all role profiles with their roles."""
	profiles = frappe.get_all("Role Profile", fields=["name"])

	result = []
	for profile in profiles:
		profile_doc = frappe.get_doc("Role Profile", profile.name)
		result.append({"name": profile.name, "roles": [r.role for r in profile_doc.roles]})

	return result


def _create_export_snapshot(export_data):
	"""Create a Permission Snapshot for the export."""
	snapshot = frappe.new_doc("Permission Snapshot")
	snapshot.snapshot_type = "Selective Export"
	snapshot.description = f"Export created at {now_datetime()}"
	snapshot.snapshot_data = json.dumps(export_data)

	if export_data["meta"]["filters"]["users"]:
		snapshot.user_filter = ", ".join(export_data["meta"]["filters"]["users"])
	if export_data["meta"]["filters"]["roles"]:
		snapshot.role_filter = ", ".join(export_data["meta"]["filters"]["roles"])
	if export_data["meta"]["filters"]["doctypes"]:
		snapshot.doctype_filter = ", ".join(export_data["meta"]["filters"]["doctypes"])

	snapshot.insert(ignore_permissions=True)
	return snapshot.name


def _generate_excel(export_data):
	"""Generate Excel file from export data."""
	try:
		import openpyxl
		from openpyxl.styles import Font, PatternFill
	except ImportError:
		frappe.throw(_("openpyxl is required for Excel export. Install with: pip install openpyxl"))

	wb = openpyxl.Workbook()

	# Header style
	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

	# Role Assignments sheet
	ws = wb.active
	ws.title = "Role Assignments"
	headers = ["User", "Role"]
	for col, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col, value=header)
		cell.font = header_font
		cell.fill = header_fill

	row = 2
	for assignment in export_data.get("role_assignments", []):
		for role in assignment.get("roles", []):
			ws.cell(row=row, column=1, value=assignment["user"])
			ws.cell(row=row, column=2, value=role)
			row += 1

	# User Permissions sheet
	ws = wb.create_sheet("User Permissions")
	headers = ["User", "Allow DocType", "For Value", "Applicable For", "Apply To All DocTypes"]
	for col, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col, value=header)
		cell.font = header_font
		cell.fill = header_fill

	row = 2
	for perm in export_data.get("user_permissions", []):
		ws.cell(row=row, column=1, value=perm.get("user"))
		ws.cell(row=row, column=2, value=perm.get("allow"))
		ws.cell(row=row, column=3, value=perm.get("for_value"))
		ws.cell(row=row, column=4, value=perm.get("applicable_for"))
		ws.cell(row=row, column=5, value="Yes" if perm.get("apply_to_all_doctypes") else "No")
		row += 1

	# Role Permissions sheet
	ws = wb.create_sheet("Role Permissions")
	headers = ["Role", "DocType", "Read", "Write", "Create", "Delete", "Submit", "Cancel", "Amend", "Report", "Export", "Import", "Share", "Print", "Email"]
	for col, header in enumerate(headers, 1):
		cell = ws.cell(row=1, column=col, value=header)
		cell.font = header_font
		cell.fill = header_fill

	row = 2
	for perm in export_data.get("role_permissions", []):
		ws.cell(row=row, column=1, value=perm.get("role"))
		ws.cell(row=row, column=2, value=perm.get("doctype"))
		ws.cell(row=row, column=3, value=perm.get("read", 0))
		ws.cell(row=row, column=4, value=perm.get("write", 0))
		ws.cell(row=row, column=5, value=perm.get("create", 0))
		ws.cell(row=row, column=6, value=perm.get("delete", 0))
		ws.cell(row=row, column=7, value=perm.get("submit", 0))
		ws.cell(row=row, column=8, value=perm.get("cancel", 0))
		ws.cell(row=row, column=9, value=perm.get("amend", 0))
		ws.cell(row=row, column=10, value=perm.get("report", 0))
		ws.cell(row=row, column=11, value=perm.get("export", 0))
		ws.cell(row=row, column=12, value=perm.get("import", 0))
		ws.cell(row=row, column=13, value=perm.get("share", 0))
		ws.cell(row=row, column=14, value=perm.get("print", 0))
		ws.cell(row=row, column=15, value=perm.get("email", 0))
		row += 1

	# Save to file
	from io import BytesIO

	file_buffer = BytesIO()
	wb.save(file_buffer)
	file_buffer.seek(0)

	# Create file in Frappe
	filename = f"permission_export_{frappe.utils.now().replace(' ', '_').replace(':', '-')}.xlsx"
	file_doc = frappe.get_doc(
		{
			"doctype": "File",
			"file_name": filename,
			"content": file_buffer.getvalue(),
			"is_private": 0,
		}
	)
	file_doc.insert(ignore_permissions=True)

	return file_doc.file_url


def _to_bool(value):
	"""Convert various truthy values to boolean."""
	if isinstance(value, bool):
		return value
	if isinstance(value, str):
		return value.lower() in ("true", "1", "yes")
	return bool(value)
